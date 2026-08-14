"""The MCP server scaffold (#260): meta tool bodies, registration, and the entry point.

``@app.tool`` returns the original plain function, so ``health`` / ``version`` are called
directly here — no fake MCP client needed. ``main()`` is exercised with ``app.run``
patched out: the real call blocks forever serving the stdio protocol loop.
"""

import asyncio
import io
import math
from datetime import date
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pandas as pd
import pytest
import quality_core.spc.capability as capability
from fastmcp.exceptions import ToolError
from mcp_app import __version__
from mcp_app.server import (
    app,
    controlplan_build,
    controlplan_build_from_project,
    controlplan_recommend_chart,
    controlplan_source_index,
    export_csv,
    fmea_chart_heatmap_png,
    fmea_chart_pareto_png,
    fmea_export_excel,
    fmea_export_pdf,
    fmea_get_scale,
    fmea_list_scales,
    fmea_run,
    fmea_run_relational,
    fmea_score,
    health,
    main,
    msa_export_excel,
    msa_export_pdf,
    msa_export_results_csv,
    msa_export_study_csv,
    msa_gage_rr,
    spc_apply_imr,
    spc_apply_xbar_r,
    spc_apply_xbar_s,
    spc_assess_stability,
    spc_c,
    spc_capability,
    spc_config_from_project,
    spc_cusum,
    spc_detect_nelson_violations,
    spc_detect_we_violations,
    spc_ewma,
    spc_export_capability_excel,
    spc_export_capability_pdf,
    spc_export_control_chart_excel,
    spc_export_control_chart_pdf,
    spc_freeze_imr,
    spc_freeze_xbar_r,
    spc_freeze_xbar_s,
    spc_imr,
    spc_normality_test,
    spc_p,
    spc_u,
    spc_xbar_r,
    spc_xbar_s,
    version,
)
from msa_app.schema import load_gage_study_csv
from quality_core.schema import (
    Action,
    ActionStatus,
    FMEADataset,
    FMEARow,
    flat_to_relational,
)

from fmea_app.rpn_engine import ACTION_COLUMNS, dataframe_to_relational

# Reused verbatim from apps/fmea/tests/test_relational_pipeline.py: 3 rows, 2 functions,
# a shared (deduplicated) effect, non-monotonic IDs so ranking + adapter sort both matter.
_ROWS = [
    dict(ID=3, Process_Step="Mix", Component="Resin", Function="Bond layers",
         Failure_Mode="Uncured", Effect="Delamination", Severity=9,
         Cause="Low temperature", Occurrence=8, Current_Control="Oven", Detection=5),
    dict(ID=1, Process_Step="Seal", Component="Edge", Function="Seal edge",
         Failure_Mode="Void", Effect="Leak", Severity=6,
         Cause="Gap", Occurrence=3, Current_Control="Visual", Detection=7),
    dict(ID=2, Process_Step="Mix", Component="Resin", Function="Bond layers",
         Failure_Mode="Uncured", Effect="Delamination", Severity=9,
         Cause="Contamination", Occurrence=2, Current_Control="Oven", Detection=5),
]


def _relational_model_dict() -> dict:
    """The plain-dict form of the _ROWS relational model a JSON-RPC client would send."""
    df = pd.DataFrame(_ROWS)
    return dataframe_to_relational(df).model_dump()


def test_health_reports_ok():
    assert health() == {"status": "ok"}


def test_version_reports_package_version():
    assert version() == {"version": __version__}


def test_exactly_the_expected_tools_are_registered():
    # Proves the decorators registered every tool on the FastMCP app object, and that
    # nothing else crept in — a stray tool would otherwise go unnoticed until a later
    # M1 issue. `list_tools` is async in fastmcp 3.4.6, hence asyncio.run.
    # Naming convention (#260): meta tools unprefixed, domain tools `<domain>_`.
    tools = asyncio.run(app.list_tools())
    assert {tool.name for tool in tools} == {
        "health",
        "version",
        "fmea_score",
        "fmea_run",
        "fmea_run_relational",
        "fmea_list_scales",
        "fmea_get_scale",
        "spc_xbar_r",
        "spc_xbar_s",
        "spc_imr",
        "spc_p",
        "spc_c",
        "spc_u",
        "spc_ewma",
        "spc_cusum",
        "spc_freeze_xbar_r",
        "spc_freeze_xbar_s",
        "spc_freeze_imr",
        "spc_apply_xbar_r",
        "spc_apply_xbar_s",
        "spc_apply_imr",
        "spc_detect_we_violations",
        "spc_detect_nelson_violations",
        "spc_capability",
        "spc_normality_test",
        "spc_assess_stability",
        "spc_config_from_project",
        "msa_gage_rr",
        "controlplan_build",
        "controlplan_build_from_project",
        "controlplan_recommend_chart",
        "controlplan_source_index",
        "export_csv",
        "fmea_export_excel",
        "fmea_export_pdf",
        "fmea_chart_pareto_png",
        "fmea_chart_heatmap_png",
        "spc_export_control_chart_excel",
        "spc_export_control_chart_pdf",
        "spc_export_capability_excel",
        "spc_export_capability_pdf",
        "msa_export_excel",
        "msa_export_pdf",
        "msa_export_study_csv",
        "msa_export_results_csv",
    }


def test_main_runs_the_server(monkeypatch: pytest.MonkeyPatch):
    calls: list[bool] = []
    monkeypatch.setattr(app, "run", lambda: calls.append(True))
    main()
    assert calls == [True]


# ---------------------------------------------------------------------------
# fmea_score — golden RPN/AP (verified against quality_core.scoring, scale-independent)
# ---------------------------------------------------------------------------


def test_fmea_score_worst_case():
    assert fmea_score(10, 10, 10) == {"rpn": 1000, "action_priority": "High"}


def test_fmea_score_low():
    assert fmea_score(9, 1, 1) == {"rpn": 9, "action_priority": "Low"}


def test_fmea_score_medium():
    # S 7-8 band, O 4-5 band, D 5-6 column -> Medium in _AP_GRID; 8*4*5 == 160.
    assert fmea_score(8, 4, 5) == {"rpn": 160, "action_priority": "Medium"}


def test_fmea_score_is_scale_independent():
    # Loading/listing any scale must not perturb the pure-math scorers.
    fmea_list_scales()
    fmea_get_scale("2019")
    fmea_get_scale("fmea4")
    assert fmea_score(8, 4, 5) == {"rpn": 160, "action_priority": "Medium"}


# ---------------------------------------------------------------------------
# fmea_run — golden table over the flat pipeline
# ---------------------------------------------------------------------------


def test_fmea_run_golden_table_ranked_flagged_and_scored():
    rows = [
        dict(ID=1, Process_Step="A", Component="C1", Function="F1",
             Failure_Mode="M1", Effect="E1", Severity=9,
             Cause="X", Occurrence=3, Current_Control="V", Detection=4),   # RPN 108, Red
        dict(ID=2, Process_Step="B", Component="C2", Function="F2",
             Failure_Mode="M2", Effect="E2", Severity=2,
             Cause="Y", Occurrence=2, Current_Control="V", Detection=2),   # RPN 8, Green
        dict(ID=3, Process_Step="C", Component="C3", Function="F3",
             Failure_Mode="M3", Effect="E3", Severity=3,
             Cause="Z", Occurrence=5, Current_Control="V", Detection=5),   # RPN 75, Yellow
    ]
    out = fmea_run(rows)
    # RPN descending -> IDs [1, 3, 2]
    assert [r["ID"] for r in out] == [1, 3, 2]
    assert [r["RPN"] for r in out] == [108, 75, 8]
    # AP populated on every row (scalar action_priority per row)
    assert [r["AP"] for r in out] == ["Low", "Low", "Low"]
    # Risk_Tier spans all three tiers; Flag_High_Severity only for S>=9
    assert [r["Risk_Tier"] for r in out] == ["Red", "Yellow", "Green"]
    assert [r["Flag_High_Severity"] for r in out] == [True, False, False]


# ---------------------------------------------------------------------------
# fmea_run_relational — golden table + action-column variance
# ---------------------------------------------------------------------------


def test_fmea_run_relational_golden_table():
    out = fmea_run_relational(_relational_model_dict())
    assert [r["ID"] for r in out] == [3, 1, 2]
    assert [r["RPN"] for r in out] == [360, 126, 90]
    assert [r["AP"] for r in out] == ["High", "Low", "Medium"]
    # An action-free model omits ACTION_COLUMNS entirely.
    assert not any(c in out[0] for c in ACTION_COLUMNS)


def test_fmea_run_relational_action_bearing_model_adds_action_columns():
    dataset = FMEADataset(rows=[FMEARow(**r) for r in _ROWS])  # type: ignore[arg-type]
    model = flat_to_relational(dataset)
    model.functions[0].failure_modes[0].links[0].action = Action(
        owner="Quality Eng", status=ActionStatus.CLOSED, due=date(2026, 8, 1), o_after=1
    )
    out = fmea_run_relational(model.model_dump())
    assert all(c in out[0] for c in ACTION_COLUMNS)


# ---------------------------------------------------------------------------
# fmea_list_scales / fmea_get_scale happy paths (each if/elif arm)
# ---------------------------------------------------------------------------


def test_fmea_list_scales_returns_the_two_builtins():
    scales = fmea_list_scales()
    assert {s["id"] for s in scales} == {"2019", "fmea4"}
    assert all(s["name"] for s in scales)  # names come from the bundled files


def test_fmea_get_scale_default_2019():
    scale = fmea_get_scale("2019")
    assert set(scale) >= {"name", "severity", "occurrence", "detection"}


def test_fmea_get_scale_fmea4():
    scale = fmea_get_scale("fmea4")
    assert set(scale) >= {"name", "severity", "occurrence", "detection"}


def test_fmea_get_scale_custom_valid_json():
    valid = fmea_get_scale("2019")
    import json

    custom = fmea_get_scale(
        "custom",
        json.dumps(
            {
                "severity": valid["severity"],
                "occurrence": valid["occurrence"],
                "detection": valid["detection"],
            }
        ),
    )
    assert set(custom) >= {"severity", "occurrence", "detection"}


# ---------------------------------------------------------------------------
# Edge cases (spec §4) — each engine ValueError/ValidationError flows through as ToolError
# ---------------------------------------------------------------------------


def test_fmea_score_out_of_range_raises_toolerror():
    with pytest.raises(ToolError, match="Severity score 11 is out of range"):
        fmea_score(11, 1, 1)


def test_fmea_score_out_of_range_carries_endash_range():
    # The engine's own message uses an en-dash; _call must not rewrite it.
    with pytest.raises(ToolError, match="Valid range is 1–10"):
        fmea_score(0, 1, 1)


def test_fmea_run_empty_rows_raises_toolerror():
    with pytest.raises(ToolError, match="Input DataFrame is empty"):
        fmea_run([])


def test_fmea_run_missing_column_raises_toolerror():
    with pytest.raises(ToolError, match="Missing required column"):
        fmea_run([{"ID": 1}])


def test_fmea_run_relational_invalid_model_raises_toolerror():
    # Structurally invalid model -> pydantic ValidationError arm of _call.
    with pytest.raises(ToolError, match="validation error"):
        fmea_run_relational({"functions": [{"id": "f1"}]})


def test_fmea_get_scale_custom_without_json_raises_toolerror():
    with pytest.raises(ToolError, match="requires custom_json"):
        fmea_get_scale("custom", None)


def test_fmea_get_scale_custom_bad_json_raises_toolerror():
    with pytest.raises(ToolError, match="Could not parse rating-scale JSON"):
        fmea_get_scale("custom", "not json")


def test_fmea_get_scale_unknown_id_raises_toolerror():
    with pytest.raises(ToolError, match="Unknown scale_id 'bogus'"):
        fmea_get_scale("bogus")


# ===========================================================================
# SPC tools (#263). Golden datasets and expected numbers are lifted verbatim
# from packages/quality-core/tests/test_spc_*.py — the MCP layer adds no
# arithmetic, so the boundary must return the already-verified engine numbers.
# ===========================================================================

XBAR_R_SAMPLE = [
    [10, 11, 12, 13, 14],
    [11, 12, 13, 14, 15],
    [9, 10, 11, 12, 13],
]
XBAR_S_SAMPLE = [list(range(1, 13)), list(range(2, 14)), list(range(3, 15))]
IMR_SAMPLE = [10, 12, 11, 15, 14]
P_COUNTS = [3, 5, 4]
P_SAMPLE_SIZES = [100, 120, 80]
C_COUNTS = [4, 7, 5, 6]
U_COUNTS = [2, 4, 3]
U_SAMPLE_SIZES = [1.0, 2.0, 1.5]

# Phase II data (test_spc_control_charts.py:282-292) — different means/ranges
# from the baseline, so "limits unchanged" is a real claim, not a tautology.
XBAR_R_PHASE_II_DATA = [[20, 21, 22, 23, 24], [19, 18, 17, 16, 15]]
XBAR_S_PHASE_II_DATA = [list(range(20, 32)), list(range(30, 18, -1))]
IMR_PHASE_II_DATA = [50, 55, 48, 60, 52, 47]

EWMA_VALUES = [11.0, 9.5, 10.2, 10.8, 9.9]
MU0 = 10.0
SIGMA = 1.0


# ---------------------------------------------------------------------------
# Charts — golden limits + points, not merely "no error"
# ---------------------------------------------------------------------------


def test_spc_xbar_r_golden_limits_and_points():
    r = spc_xbar_r(XBAR_R_SAMPLE)
    assert r["subgroup_means"] == pytest.approx([12.0, 13.0, 11.0])
    assert r["ranges"] == pytest.approx([4.0, 4.0, 4.0])
    assert r["ucl_x"] == pytest.approx(14.308, rel=1e-4)
    assert r["lcl_x"] == pytest.approx(9.692, rel=1e-4)
    assert r["ucl_r"] == pytest.approx(8.456, rel=1e-4)
    assert r["lcl_r"] == pytest.approx(0.0)
    assert r["sigma_hat"] == pytest.approx(4.0 / 2.326, rel=1e-4)


def test_spc_xbar_s_golden_limits():
    r = spc_xbar_s(XBAR_S_SAMPLE)
    subgroup_std = math.sqrt(13.0)
    assert r["ucl_x"] == pytest.approx(7.5 + 0.886 * subgroup_std, rel=1e-4)
    assert r["sigma_hat"] == pytest.approx(subgroup_std / 0.9776, rel=1e-4)


def test_spc_imr_golden_limits_and_moving_ranges():
    r = spc_imr(IMR_SAMPLE)
    assert r["moving_ranges"] == pytest.approx([2.0, 1.0, 4.0, 1.0])
    assert r["ucl_x"] == pytest.approx(17.72, rel=1e-4)
    assert r["lcl_x"] == pytest.approx(7.08, rel=1e-4)
    assert r["ucl_mr"] == pytest.approx(6.534, rel=1e-4)
    assert r["sigma_hat"] == pytest.approx(2.0 / 1.128, rel=1e-4)


def test_spc_p_golden_pbar_and_proportions():
    r = spc_p(P_COUNTS, P_SAMPLE_SIZES)
    assert r["pbar"] == pytest.approx(12.0 / 300.0, rel=1e-4)
    assert r["proportions"] == pytest.approx([0.03, 5 / 120, 0.05], rel=1e-4)
    pbar = 12.0 / 300.0
    assert r["ucl"][0] == pytest.approx(
        pbar + 3.0 * math.sqrt((pbar * (1.0 - pbar)) / 100.0), rel=1e-4
    )
    assert r["lcl"][0] == pytest.approx(0.0)


def test_spc_c_golden_cbar_and_limits():
    r = spc_c(C_COUNTS)
    assert r["cbar"] == pytest.approx(5.5)
    assert r["ucl"] == pytest.approx(5.5 + 3.0 * math.sqrt(5.5), rel=1e-4)
    assert r["lcl"] == pytest.approx(0.0)


def test_spc_u_golden_ubar_and_limits():
    r = spc_u(U_COUNTS, U_SAMPLE_SIZES)
    assert r["ubar"] == pytest.approx(2.0)
    assert r["ucl"][1] == pytest.approx(2.0 + 3.0 * math.sqrt(2.0 / 2.0), rel=1e-4)
    assert r["lcl"][0] == pytest.approx(0.0)


def test_spc_ewma_golden_recursion_and_echoed_defaults():
    r = spc_ewma(EWMA_VALUES, MU0, SIGMA)
    # z0 seeds from mu0, not x0 (test_spc_ewma.py:41-54).
    z0 = 0.2 * EWMA_VALUES[0] + 0.8 * MU0
    assert r["z"][0] == pytest.approx(z0)
    assert r["z"][1] == pytest.approx(0.2 * EWMA_VALUES[1] + 0.8 * z0)
    # Defaults pulled from the engine constants, echoed on the result.
    assert r["lam"] == pytest.approx(0.20)
    assert r["L"] == pytest.approx(2.860)
    assert r["pairing_adequate"] is True


def test_spc_ewma_mismatched_pairing_is_flagged_not_swallowed():
    # A tabulated lambda paired with a non-tabulated L must come back flagged.
    r = spc_ewma(EWMA_VALUES, MU0, SIGMA, lam=0.05, L=3.0)
    assert r["pairing_adequate"] is False
    assert r["pairing_note"] != ""


def test_spc_ewma_untabulated_lambda_pairing_ok_branch():
    # Custom lambda not in EWMA_L_BY_LAMBDA -> the "no key matches" fallback: adequate.
    r = spc_ewma(EWMA_VALUES, MU0, SIGMA, lam=0.15, L=3.0)
    assert r["pairing_adequate"] is True
    assert r["pairing_note"] == ""


def test_spc_cusum_golden_recursion_and_echoed_defaults():
    r = spc_cusum([11.0, 9.5, 10.2], MU0, SIGMA)
    z0 = (11.0 - MU0) / SIGMA
    assert r["c_plus"][0] == pytest.approx(max(0.0, z0 - 0.5))
    assert r["c_minus"][0] == pytest.approx(max(0.0, -z0 - 0.5))
    assert r["k"] == pytest.approx(0.5)
    assert r["h"] == pytest.approx(5.0)
    assert r["fir"] is False


def test_spc_cusum_fir_head_start_seeds_both_arms():
    r = spc_cusum([MU0] * 5, MU0, SIGMA, fir=True)
    # seed = 0.5 * h = 2.5; z0 = 0 -> c_plus0 = max(0, -0.5 + 2.5) = 2.0.
    assert r["fir"] is True
    assert r["c_plus"][0] == pytest.approx(2.0)
    assert r["c_minus"][0] == pytest.approx(2.0)


# ---------------------------------------------------------------------------
# Phase — freeze + apply round trip (the acceptance criterion)
# ---------------------------------------------------------------------------


def test_spc_freeze_xbar_r_returns_frozen_limits_shape():
    frozen = spc_freeze_xbar_r(XBAR_R_SAMPLE)
    assert frozen["chart_type"] == "xbar_r"
    assert frozen["n"] == 5
    assert frozen["sigma_method"] == "Rbar/d2"
    # 3 subgroups is below the AIAG baseline floor -> soft-flagged, never raised.
    assert frozen["baseline_adequate"] is False
    assert frozen["baseline_note"] != ""


def test_spc_apply_xbar_r_round_trip_uses_frozen_limits_verbatim():
    frozen = spc_freeze_xbar_r(XBAR_R_SAMPLE)
    r = spc_apply_xbar_r(XBAR_R_PHASE_II_DATA, frozen)
    # Phase II limits equal the frozen ones, nothing recomputed from new data.
    assert r["xbarbar"] == pytest.approx(frozen["center_line"])
    assert r["rbar"] == pytest.approx(frozen["dispersion_center"])
    assert r["ucl_x"] == pytest.approx(frozen["ucl_x"])
    assert r["lcl_x"] == pytest.approx(frozen["lcl_x"])
    assert r["ucl_r"] == pytest.approx(frozen["ucl_disp"])
    assert r["lcl_r"] == pytest.approx(frozen["lcl_disp"])
    assert r["sigma_hat"] == pytest.approx(frozen["sigma_hat"])
    # But the plotted points ARE the new data's, not the baseline's.
    assert r["subgroup_means"] == pytest.approx([22.0, 17.0])
    assert r["subgroup_means"] != pytest.approx(
        spc_xbar_r(XBAR_R_SAMPLE)["subgroup_means"]
    )


def test_spc_apply_xbar_s_round_trip_uses_frozen_limits_verbatim():
    frozen = spc_freeze_xbar_s(XBAR_S_SAMPLE)
    r = spc_apply_xbar_s(XBAR_S_PHASE_II_DATA, frozen)
    assert r["ucl_x"] == pytest.approx(frozen["ucl_x"])
    assert r["lcl_x"] == pytest.approx(frozen["lcl_x"])
    assert r["ucl_s"] == pytest.approx(frozen["ucl_disp"])
    assert r["sigma_hat"] == pytest.approx(frozen["sigma_hat"])


def test_spc_apply_imr_round_trip_uses_frozen_limits_verbatim():
    frozen = spc_freeze_imr(IMR_SAMPLE)
    r = spc_apply_imr(IMR_PHASE_II_DATA, frozen)
    assert r["xbar"] == pytest.approx(frozen["center_line"])
    assert r["ucl_x"] == pytest.approx(frozen["ucl_x"])
    assert r["lcl_x"] == pytest.approx(frozen["lcl_x"])
    assert r["ucl_mr"] == pytest.approx(frozen["ucl_disp"])
    assert r["sigma_hat"] == pytest.approx(frozen["sigma_hat"])
    assert r["values"] == pytest.approx(IMR_PHASE_II_DATA)


def test_spc_freeze_xbar_r_exclusion_shifts_limits():
    baseline = [list(row) for row in XBAR_R_SAMPLE] + [[10, 100, 5, 95, 50]]
    no_excl = spc_freeze_xbar_r(baseline)
    with_excl = spc_freeze_xbar_r(
        baseline, excluded=[{"index": 3, "cause": "tool jam, log #4471"}]
    )
    assert with_excl["ucl_x"] != pytest.approx(no_excl["ucl_x"])


def test_spc_freeze_imr_phase_i_range_recorded_verbatim():
    # tuple[str, str] survives the boundary as a 2-tuple/array (JSON round-trip).
    rng = ("2026-01-01T00:00:00+00:00", "2026-02-01T00:00:00+00:00")
    frozen = spc_freeze_imr(IMR_SAMPLE, phase_i_range=rng)
    assert tuple(frozen["phase_i_range"]) == rng


# ---------------------------------------------------------------------------
# Rules — WE + Nelson flagged at the right indices on golden series
# ---------------------------------------------------------------------------


def test_spc_detect_we_violations_golden_indices():
    # A >3sigma point fires WE Rule 1 at index 2 (test_spc_rule_detection.py:31).
    v = spc_detect_we_violations([0.1, 0.2, 3.2], cl=0.0, sigma=1.0)
    assert {"index": 2, "rule": "Western Electric Rule 1"} in v


def test_spc_detect_we_violations_clean_series_is_empty():
    assert spc_detect_we_violations([0.2, -0.1, 0.4, -0.3, 0.1], cl=0.0, sigma=1.0) == []


def test_spc_detect_nelson_violations_golden_run_rule():
    # 9-in-a-row fires Nelson Rule 2 at index 8 (test_spc_rule_detection.py:283).
    v = spc_detect_nelson_violations([1.0] * 9, cl=0.0, sigma=1.0)
    assert v == [{"index": 8, "rule": "Nelson Rule 2"}]


def test_spc_detect_nelson_eight_in_a_row_does_not_fire_run_rule():
    # Nelson's run test is 9-in-a-row, not WE's 8 (boundary, one below the edge).
    assert spc_detect_nelson_violations([1.0] * 8, cl=0.0, sigma=1.0) == []


# ---------------------------------------------------------------------------
# Capability — indices + the #193 CI/df/estimator pairing, per method path
# ---------------------------------------------------------------------------

NORMAL_DATA = np.random.default_rng(0).normal(10.0, 0.5, size=40).tolist()


def test_spc_capability_normal_path_indices_and_ci_pairing():
    study = spc_capability(NORMAL_DATA, 8.0, 12.0, force_method="normal")
    assert study["method"] == "normal"
    for key in ("cp", "cpk", "pp", "ppk"):
        assert study[key] is not None
    # #193: the parametric CI rides Pp/Ppk (ddof=1), never the within-sigma Cp/Cpk.
    assert study["cp_ci"] is None
    assert study["cpk_ci"] is None
    assert study["pp_ci"] is not None
    assert study["ppk_ci"] is not None
    assert study["ci_estimator"] == "sample_sd_ddof1"
    assert study["ci_df"] == study["n"] - 1


def test_spc_capability_percentile_path_flips_the_ci_pairing(monkeypatch):
    monkeypatch.setattr(capability, "BOOTSTRAP_RESAMPLES", 50)
    study = spc_capability(NORMAL_DATA, 8.0, 12.0, force_method="percentile")
    assert study["method"] == "percentile"
    # Percentile method computes no Pp/Ppk; the bootstrap CI rides Cp/Cpk.
    assert study["pp"] is None
    assert study["ppk"] is None
    assert study["pp_ci"] is None
    assert study["ppk_ci"] is None
    assert study["cp_ci"] is not None
    assert study["cpk_ci"] is not None
    assert study["ci_estimator"] == "bootstrap_percentile"
    assert study["ci_df"] is None


def test_spc_capability_no_spec_limits_returns_null_indices_without_raising():
    study = spc_capability(NORMAL_DATA, None, None, force_method="normal")
    assert study["cp"] is None
    assert study["cpk"] is None


def test_spc_capability_violations_none_is_stable_none_not_assessed():
    # #191 D2: the `is None` distinction is load-bearing — omitted means "not assessed".
    study = spc_capability(NORMAL_DATA, 8.0, 12.0, force_method="normal")
    assert study["stable"] is None


def test_spc_capability_violations_empty_is_stable_true_in_control():
    study = spc_capability(NORMAL_DATA, 8.0, 12.0, force_method="normal", violations=[])
    assert study["stable"] is True


def test_spc_capability_violations_non_empty_is_stable_false():
    study = spc_capability(
        NORMAL_DATA,
        8.0,
        12.0,
        force_method="normal",
        violations=[{"index": 9, "rule": "Western Electric Rule 1"}],
    )
    assert study["stable"] is False
    assert study["stability_note"]


def test_spc_normality_test_golden_keys():
    r = spc_normality_test(NORMAL_DATA)
    assert set(r) >= {"w_stat", "p_value", "is_normal"}
    assert r["is_normal"] is True


# ---------------------------------------------------------------------------
# Stability — each chart_type path, long-format parallel lists
# ---------------------------------------------------------------------------


def _long_format(subgroups: list[list[float]]) -> tuple[list[float], list[int]]:
    values: list[float] = []
    labels: list[int] = []
    for index, group in enumerate(subgroups, start=1):
        for value in group:
            values.append(float(value))
            labels.append(index)
    return values, labels


def test_spc_assess_stability_imr_in_control_has_no_signals():
    values = [1.0, 2.0] * 5
    r = spc_assess_stability(values, list(range(1, 11)), chart_type="I-MR")
    assert r["sigma_hat"] > 0
    assert r["signals"] == []


def test_spc_assess_stability_imr_out_of_control_flags_signal():
    values = [1.0, 2.0, 1.0, 2.0, 1.0, 2.0, 1.0, 2.0, 1.0, 20.0]
    r = spc_assess_stability(values, list(range(1, 11)), chart_type="I-MR")
    assert len(r["signals"]) >= 1


def test_spc_assess_stability_xbar_r_path():
    values, labels = _long_format(XBAR_R_SAMPLE)
    r = spc_assess_stability(values, labels, chart_type="Xbar-R")
    assert r["sigma_hat"] > 0
    assert isinstance(r["signals"], list)


def test_spc_assess_stability_xbar_s_path():
    values, labels = _long_format(XBAR_S_SAMPLE)
    r = spc_assess_stability(values, labels, chart_type="Xbar-S")
    assert r["sigma_hat"] > 0
    assert isinstance(r["signals"], list)


# ---------------------------------------------------------------------------
# Edge / error cases — every enumerated ValueError branch surfaces as ToolError
# via the _call wrapper, matched on the exact engine message.
# ---------------------------------------------------------------------------


def test_spc_xbar_r_ragged_subgroups_raise_toolerror():
    with pytest.raises(ToolError):
        spc_xbar_r([[1.0, 2.0, 3.0], [4.0, 5.0]])


def test_spc_xbar_r_untabulated_subgroup_size_raises_toolerror():
    with pytest.raises(ToolError, match="subgroup size between 2 and 10"):
        spc_xbar_r([[1], [2], [3]])


def test_spc_xbar_s_untabulated_subgroup_size_raises_toolerror():
    with pytest.raises(ToolError, match="subgroup size between 2 and 12"):
        spc_xbar_s([[1], [2]])


def test_spc_imr_too_few_values_raises_toolerror():
    with pytest.raises(ToolError, match="at least two values"):
        spc_imr([5])


def test_spc_c_empty_raises_toolerror():
    with pytest.raises(ToolError, match="at least one count"):
        spc_c([])


def test_spc_p_mismatched_lengths_raise_toolerror():
    with pytest.raises(ToolError, match="matching 1D arrays"):
        spc_p([1, 2], [100])


def test_spc_p_non_finite_sample_size_raises_toolerror():
    # #200 regression: a NaN size must raise, not silently produce NaN limits.
    with pytest.raises(ToolError, match="sample sizes must be positive"):
        spc_p([1, 2], [10.0, float("nan")])


def test_spc_u_non_finite_sample_size_raises_toolerror():
    with pytest.raises(ToolError, match="sample sizes must be positive"):
        spc_u([1, 2], [10.0, float("inf")])


def test_spc_ewma_nonpositive_sigma_raises_toolerror():
    with pytest.raises(ToolError, match="requires sigma > 0"):
        spc_ewma(EWMA_VALUES, MU0, 0.0)


def test_spc_ewma_lambda_out_of_range_raises_toolerror():
    with pytest.raises(ToolError, match=r"requires 0 < lam <= 1"):
        spc_ewma(EWMA_VALUES, MU0, SIGMA, lam=1.5)


def test_spc_ewma_nonpositive_l_raises_toolerror():
    with pytest.raises(ToolError, match="requires L > 0"):
        spc_ewma(EWMA_VALUES, MU0, SIGMA, L=0.0)


def test_spc_cusum_nonpositive_sigma_raises_toolerror():
    with pytest.raises(ToolError, match="requires sigma > 0"):
        spc_cusum(EWMA_VALUES, MU0, 0.0)


def test_spc_cusum_nonpositive_k_raises_toolerror():
    with pytest.raises(ToolError, match="requires k > 0"):
        spc_cusum(EWMA_VALUES, MU0, SIGMA, k=0.0)


def test_spc_cusum_nonpositive_h_raises_toolerror():
    with pytest.raises(ToolError, match="requires h > 0"):
        spc_cusum(EWMA_VALUES, MU0, SIGMA, h=0.0)


def test_spc_freeze_empty_cause_raises_toolerror():
    with pytest.raises(ToolError, match="non-empty documented cause"):
        spc_freeze_xbar_r(XBAR_R_SAMPLE, excluded=[{"index": 0, "cause": "   "}])


def test_spc_freeze_out_of_range_index_raises_toolerror():
    with pytest.raises(ToolError, match="out of range for the baseline"):
        spc_freeze_xbar_r(XBAR_R_SAMPLE, excluded=[{"index": 999, "cause": "bad"}])


def test_spc_freeze_duplicate_index_raises_toolerror():
    with pytest.raises(ToolError, match="is duplicated"):
        spc_freeze_xbar_r(
            XBAR_R_SAMPLE,
            excluded=[{"index": 0, "cause": "a"}, {"index": 0, "cause": "b"}],
        )


def test_spc_apply_chart_type_mismatch_raises_toolerror():
    # An I-MR baseline fed into the X-bar & R Phase II tool.
    imr_frozen = spc_freeze_imr(IMR_SAMPLE)
    with pytest.raises(ToolError, match="expected 'xbar_r'"):
        spc_apply_xbar_r(XBAR_R_PHASE_II_DATA, imr_frozen)


def test_spc_apply_subgroup_size_mismatch_raises_toolerror():
    frozen = spc_freeze_xbar_r(XBAR_R_SAMPLE)  # n=5
    with pytest.raises(ToolError, match="computed for n=5"):
        spc_apply_xbar_r([[1, 2, 3], [4, 5, 6]], frozen)


def test_spc_detect_we_violations_nonpositive_sigma_raises_toolerror():
    with pytest.raises(ToolError, match="sigma must be positive"):
        spc_detect_we_violations([0.1, 0.2, 0.3], cl=0.0, sigma=0.0)


def test_spc_detect_nelson_violations_nonpositive_sigma_raises_toolerror():
    with pytest.raises(ToolError, match="sigma must be positive"):
        spc_detect_nelson_violations([0.1, 0.2, 0.3], cl=0.0, sigma=-1.0)


def test_spc_capability_alpha_out_of_range_raises_toolerror():
    with pytest.raises(ToolError, match=r"alpha must be in \(0, 1\)"):
        spc_capability(NORMAL_DATA, 8.0, 12.0, alpha=1.5)


def test_spc_capability_too_few_points_raises_toolerror():
    with pytest.raises(ToolError, match="at least three observations"):
        spc_capability([1.0, 2.0], 8.0, 12.0)


def test_spc_capability_invalid_ndim_raises_toolerror():
    with pytest.raises(ToolError, match="1D individuals or 2D subgroups"):
        spc_capability([[[1.0, 2.0], [3.0, 4.0]]], 8.0, 12.0)  # type: ignore[list-item]


def test_spc_capability_constant_data_raises_toolerror():
    with pytest.raises(ToolError, match="constant data has no capability"):
        spc_capability([5.0] * 10, 1.0, 9.0)


def test_spc_normality_test_too_few_points_raises_toolerror():
    with pytest.raises(ToolError, match="at least three values"):
        spc_normality_test([10.0, 10.1])


def test_spc_assess_stability_ragged_subgroups_raise_toolerror():
    values, labels = _long_format([[1.0, 2.0, 3.0], [4.0, 5.0]])
    with pytest.raises(ToolError):
        spc_assess_stability(values, labels, chart_type="Xbar-R")


def test_spc_apply_malformed_frozen_raises_keyerror_not_toolerror():
    # KNOWN, spec-sanctioned (changes.md "For the Tester"): a structurally malformed
    # `frozen` dict raises KeyError, NOT a structured ToolError, because `_call`
    # catches only ValueError/ValidationError. Pinned here so a future change that
    # silently widens `_call` to swallow KeyError is a visible, reviewed decision.
    with pytest.raises(KeyError):
        spc_apply_xbar_r(XBAR_R_PHASE_II_DATA, {"not": "a frozen dict"})


# ===========================================================================
# Control Plan tools (#265). The three tools are thin passthroughs over
# controlplan_app.connector; the connector's own arithmetic/ordering is proved
# 100% in apps/controlplan/tests/test_connector.py. These tests pin the
# boundary: the tools return the connector's already-verified output verbatim
# (JSON-friendly), and bad input surfaces as a structured ToolError.
#
# The _ROWS / _relational_model_dict() fixture reused here is the same one
# fmea_run_relational's golden test uses (RPN 360/126/90, ID order [3,1,2]).
# _ROWS has one FailureMode per (Resin/Uncured) and (Edge/Void) — Resin/Uncured
# carries two links (row 3, O=8 and row 2, O=2), so build_control_plan yields
# exactly TWO rows (one per FailureMode), highest-risk (Resin/Uncured, High AP,
# worst link RPN 360) first.
# ===========================================================================


def _collision_model_dict() -> dict:
    """A relational model whose two FailureModes collide on component+description.

    Mirrors apps/controlplan/tests/test_connector.py::
    test_characteristic_collision_falls_back_to_failure_mode_id in flat->relational
    form: two functions share Component 'Bracket' and Failure_Mode 'Incomplete weld',
    so the base characteristic collides and the second row takes the ' (F2-M1)' suffix.
    """
    rows = [
        dict(ID=1, Process_Step="Weld", Component="Bracket", Function="Weld joint",
             Failure_Mode="Incomplete weld", Effect="Joint fails", Severity=9,
             Cause="Contamination", Occurrence=6, Current_Control="Visual", Detection=1),
        dict(ID=2, Process_Step="Rework", Component="Bracket", Function="Rework joint",
             Failure_Mode="Incomplete weld", Effect="Joint fails", Severity=3,
             Cause="Operator error", Occurrence=3, Current_Control="Visual", Detection=3),
    ]
    return dataframe_to_relational(pd.DataFrame(rows)).model_dump()


# ---------------------------------------------------------------------------
# controlplan_build — golden / empty / invalid
# ---------------------------------------------------------------------------


def test_controlplan_build_golden_rows():
    rows = controlplan_build(_relational_model_dict())
    # One row per FailureMode: (Resin/Uncured) and (Edge/Void) -> 2 rows.
    assert len(rows) == 2
    # Highest-AP/RPN row first: Resin/Uncured is High AP (worst link RPN 360),
    # Edge/Void is Low AP (RPN 126).
    assert [r["characteristic"] for r in rows] == ["Resin — Uncured", "Edge — Void"]
    # Every row flags its sample plan as a connector placeholder (F-10, #196)...
    assert all(r["sample_plan_is_placeholder"] is True for r in rows)
    # ...and recommended_chart is always null from build (no data-type/n input).
    assert all(r["recommended_chart"] is None for r in rows)
    # measurement_method comes from the worst link's control; characteristic populated.
    assert [r["measurement_method"] for r in rows] == ["Oven", "Visual"]
    assert all(r["characteristic"] for r in rows)


# ===========================================================================
# Export / report / PNG tools (M1-7, #266). These prove ONLY the thin-wrapper
# contract: valid input -> the right File/Image with the right magic bytes and
# mime type; sanitization is inherited (not bypassed); empty input is a
# structured ToolError; a missing report key/column surfaces per the coder's
# documented KeyError policy. The exporters'/sanitizer's own internals have
# their own 100%-covered suites and are not re-tested here.
# ===========================================================================

# FMEA rows straight from fmea_run's own output — Failure_Mode, RPN, Risk_Tier,
# Severity, Occurrence are all present, so the FMEA export + PNG tools get exactly
# the shape they document. Computed once; the tools never mutate it in place.
_FMEA_SCORED = fmea_run(_ROWS)

# SPC control-chart report params — a chart tool's own result fields passed back in
# (lifted from apps/spc/tests/test_exporter.py's fixture, the shape the builder takes).
_CC_KW = dict(
    chart_label="Xbar-R Chart",
    stream="ply_thickness",
    rule_set="Western Electric",
    points=[10.0, 10.5, 13.9, 9.8, 10.1],
    cl=10.0,
    ucl=12.0,
    lcl=8.0,
    violations=[{"index": 2, "rule": "Rule 1: beyond 3-sigma"}],
    metrics=[("Xbarbar", "10.0000"), ("Rbar", "1.2000")],
)

# SPC capability report params — spc_capability's / spc_normality_test's own dicts verbatim.
_CAP_KW = dict(
    stream_label="Ply Thickness",
    values=[10.0, 10.1, 9.9, 10.2, 9.8, 10.05],
    capability={
        "cp": 1.45, "cpk": 1.21, "pp": 1.40, "ppk": 1.18,
        "mean": 10.0, "sigma_hat": 0.05, "sigma_overall": 0.06,
    },
    lsl=9.7,
    usl=10.3,
    normality={"w_stat": 0.98, "p_value": 0.61, "is_normal": True},
    oos_signal_count=0,
)

# MSA study rows + compute_gage_rr's result dict verbatim (from apps/msa/tests/test_exporter.py).
_MSA_STUDY = [
    {"part": "P01", "appraiser": "A", "trial": 1, "measurement": 10.05},
    {"part": "P01", "appraiser": "A", "trial": 2, "measurement": 10.02},
    {"part": "P02", "appraiser": "B", "trial": 1, "measurement": 9.98},
    {"part": "P02", "appraiser": "B", "trial": 2, "measurement": 10.01},
]
_MSA_RESULTS = {
    "ev": 0.03, "av": 0.02, "grr": 0.036, "pv": 0.5, "tv": 0.501,
    "pev_study": 5.99, "pav_study": 3.99, "pgrr_study": 7.19, "ppv_study": 99.80,
    "pev_tolerance": 10.0, "pav_tolerance": 8.0, "pgrr_tolerance": 12.0,
    "ppv_tolerance": 150.0, "ndc": 6, "verdict": "Accept", "mean": 10.0,
    "n_parts": 2, "n_appraisers": 2, "n_trials": 2, "is_balanced": True,
    "method": "average_and_range",
    "method_note": "Average-and-Range method: the part x appraiser interaction is NOT estimated.",
}

_PNG_MAGIC = b"\x89PNG\r\n\x1a\n"


def _resource_mime(file_result) -> str:
    return file_result.to_resource_content().resource.mimeType


# ---------------------------------------------------------------------------
# Happy paths — right File/Image, right magic bytes, right mime type
# ---------------------------------------------------------------------------


def test_export_csv_happy_path_returns_csv_file():
    result = export_csv(_FMEA_SCORED)
    text = result.data.decode("utf-8")
    header = text.splitlines()[0]
    # header carries the FMEA columns, and the highest-RPN failure mode is in the body.
    assert "Failure_Mode" in header and "RPN" in header
    assert "Uncured" in text
    assert _resource_mime(result) == "application/csv"


def test_fmea_export_excel_happy_path_is_xlsx():
    result = fmea_export_excel(_FMEA_SCORED)
    assert result.data[:2] == b"PK"  # zip/OOXML magic
    assert _resource_mime(result) == "application/xlsx"


def test_fmea_export_pdf_happy_path_is_pdf():
    result = fmea_export_pdf(_FMEA_SCORED)
    assert result.data[:4] == b"%PDF"
    assert _resource_mime(result) == "application/pdf"


def test_fmea_chart_pareto_png_happy_path_is_png():
    result = fmea_chart_pareto_png(_FMEA_SCORED)
    assert result.data[:8] == _PNG_MAGIC
    assert result.to_image_content().mimeType == "image/png"


def test_fmea_chart_heatmap_png_happy_path_is_png():
    result = fmea_chart_heatmap_png(_FMEA_SCORED)
    assert result.data[:8] == _PNG_MAGIC
    assert result.to_image_content().mimeType == "image/png"


def test_spc_export_control_chart_excel_happy_path_is_xlsx():
    result = spc_export_control_chart_excel(**_CC_KW)
    assert result.data[:2] == b"PK"
    assert _resource_mime(result) == "application/xlsx"


def test_spc_export_control_chart_pdf_happy_path_is_pdf():
    result = spc_export_control_chart_pdf(**_CC_KW)
    assert result.data[:4] == b"%PDF"
    assert _resource_mime(result) == "application/pdf"


def test_spc_export_control_chart_excel_with_secondary_series_renders_it():
    # Both secondary_label AND secondary_points supplied -> the report's optional
    # second column is rendered (the True arm of _control_chart_report's guard).
    wb = openpyxl.load_workbook(
        io.BytesIO(
            spc_export_control_chart_excel(
                **_CC_KW, secondary_label="C-", secondary_points=[0.0, 0.1, 0.2, 0.0, 0.0]
            ).data
        )
    )
    # The per-point sheet gains a column named after the secondary label.
    header = [c.value for c in next(wb.worksheets[0].iter_rows(max_row=1))]
    assert "C-" in header


def test_spc_export_control_chart_excel_half_a_secondary_pair_is_ignored():
    # Only the label given (points None) -> guard's second operand is False; no second
    # column. Proves the "both or neither" branch, distinct from the default (label None).
    wb = openpyxl.load_workbook(
        io.BytesIO(spc_export_control_chart_excel(**_CC_KW, secondary_label="C-").data)
    )
    header = [c.value for c in next(wb.worksheets[0].iter_rows(max_row=1))]
    assert "C-" not in header


def test_spc_export_capability_excel_happy_path_is_xlsx():
    result = spc_export_capability_excel(**_CAP_KW)
    assert result.data[:2] == b"PK"
    assert _resource_mime(result) == "application/xlsx"


def test_spc_export_capability_pdf_happy_path_is_pdf():
    result = spc_export_capability_pdf(**_CAP_KW)
    assert result.data[:4] == b"%PDF"
    assert _resource_mime(result) == "application/pdf"


def test_msa_export_excel_happy_path_is_xlsx():
    result = msa_export_excel(_MSA_STUDY, _MSA_RESULTS, usl=10.5, lsl=9.5)
    assert result.data[:2] == b"PK"
    assert _resource_mime(result) == "application/xlsx"


def test_msa_export_pdf_happy_path_is_pdf():
    result = msa_export_pdf(_MSA_STUDY, _MSA_RESULTS)
    assert result.data[:4] == b"%PDF"
    assert _resource_mime(result) == "application/pdf"


def test_msa_export_study_csv_happy_path_returns_csv():
    result = msa_export_study_csv(_MSA_STUDY)
    text = result.data.decode("utf-8")
    assert text.splitlines()[0] == "part,appraiser,trial,measurement"
    assert "P01" in text
    assert _resource_mime(result) == "application/csv"


def test_msa_export_results_csv_happy_path_returns_csv():
    result = msa_export_results_csv(_MSA_RESULTS)
    text = result.data.decode("utf-8")
    assert "Verdict" in text.splitlines()[0]
    assert "Accept" in text
    assert _resource_mime(result) == "application/csv"


# ---------------------------------------------------------------------------
# Sanitization enforcement — the issue's explicit acceptance criterion.
# The negative control that proves these load-bearing lives in test-results.md
# (server.py mutated to bypass core_export_csv; both tests below then fail).
# ---------------------------------------------------------------------------


def test_export_csv_escapes_formula_injection_and_leaves_numbers_alone():
    result = export_csv(
        [{"note": "=cmd|'/bin/calc'", "plus": "+SUM(A1)", "minus": "-1+1",
          "at": "@import", "num": "-3.0000"}]
    )
    text = result.data.decode("utf-8")
    # Every formula-leading cell comes back apostrophe-escaped.
    assert "'=cmd|'/bin/calc'" in text
    assert "'+SUM(A1)" in text
    assert "'-1+1" in text
    assert "'@import" in text
    # ...and no bare formula lead survives (strip the escaped occurrences first).
    stripped = text.replace("'=", "").replace("'+", "").replace("'@", "").replace("'-", "")
    for lead in ("=cmd", "+SUM", "@import"):
        assert lead not in stripped
    # A numeric literal must NOT be escaped (proves it doesn't over-escape).
    assert "-3.0000" in text
    assert "'-3.0000" not in text


def test_msa_export_excel_escapes_formula_injection_in_study_cell():
    # One Excel-producing tool must also escape (issue: CSV AND at least one XLSX tool).
    study = [{"part": "=cmd|'/bin/calc'", "appraiser": "A", "trial": 1, "measurement": 10.0}]
    wb = openpyxl.load_workbook(io.BytesIO(msa_export_excel(study, _MSA_RESULTS).data))
    cell = str(wb["Study Data"].cell(row=2, column=1).value)
    assert cell.startswith("'=")
    assert not cell.startswith("=cmd")  # never an unescaped leading '='


# ---------------------------------------------------------------------------
# Error paths — empty input -> structured ToolError before the exporter runs.
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "tool",
    [
        export_csv,
        fmea_export_excel,
        fmea_export_pdf,
        fmea_chart_pareto_png,
        fmea_chart_heatmap_png,
        msa_export_study_csv,
    ],
)
def test_single_list_export_tool_rejects_empty_input(tool):
    with pytest.raises(ToolError, match="at least one row"):
        tool([])


def test_msa_export_excel_rejects_empty_study():
    with pytest.raises(ToolError, match="at least one row"):
        msa_export_excel([], _MSA_RESULTS)


def test_msa_export_pdf_rejects_empty_study():
    with pytest.raises(ToolError, match="at least one row"):
        msa_export_pdf([], _MSA_RESULTS)


# ---------------------------------------------------------------------------
# controlplan_build — empty model + invalid model
# ---------------------------------------------------------------------------


def test_controlplan_build_empty_fmea_returns_empty_list():
    assert controlplan_build({"functions": []}) == []


def test_controlplan_build_invalid_model_raises_toolerror():
    with pytest.raises(ToolError, match="validation error"):
        controlplan_build({"functions": [{"id": "f1"}]})


# ---------------------------------------------------------------------------
# controlplan_build_from_project — project-file arrow (M3-2, #277)
# ---------------------------------------------------------------------------

_FIXTURE_PROJECT = (
    Path(__file__).resolve().parents[3]
    / "packages"
    / "quality-core"
    / "tests"
    / "fixtures"
    / "project"
)


def test_controlplan_build_from_project_writes_and_returns(tmp_path: Path):
    import shutil

    shutil.copytree(_FIXTURE_PROJECT / "fmea", tmp_path / "fmea")

    result = controlplan_build_from_project(str(tmp_path))

    # Written to the project-file contract path.
    assert (tmp_path / "control-plan" / "plan.json").exists()
    # Structured dict envelope + connector-derived rows (not the hand-built fixture row).
    assert result["schema_version"] == 1
    assert result["generated_by"].startswith("controlplan_app==")
    assert [r["characteristic"] for r in result["rows"]] == ["Bracket — Bore oversize"]
    assert result["rows"][0]["source_cause_id"] == "F1::F1-M1::F1-M1-C1"


def test_controlplan_build_from_project_missing_fmea_raises_toolerror(tmp_path: Path):
    # Empty project dir -> ProjectError -> structured ToolError; nothing written.
    with pytest.raises(ToolError):
        controlplan_build_from_project(str(tmp_path))
    assert not (tmp_path / "control-plan" / "plan.json").exists()


# ---------------------------------------------------------------------------
# spc_config_from_project — Control Plan → SPC project-file arrow (M3-3, #278)
# ---------------------------------------------------------------------------


def test_spc_config_from_project_writes_and_returns(tmp_path: Path):
    import shutil

    shutil.copytree(_FIXTURE_PROJECT / "control-plan", tmp_path / "control-plan")

    result = spc_config_from_project(str(tmp_path))

    # Written to the project-file contract path.
    assert (tmp_path / "spc" / "config.json").exists()
    # Structured dict envelope + config rows derived from the plan.
    assert result["schema_version"] == 1
    assert result["generated_by"].startswith("spc_app==")
    assert [r["characteristic"] for r in result["rows"]] == ["Example Characteristic"]
    assert result["rows"][0]["chart_key"] == "Xbar-R"


def test_spc_config_from_project_missing_plan_raises_toolerror(tmp_path: Path):
    # Empty project dir -> ProjectError -> structured ToolError; nothing written.
    with pytest.raises(ToolError):
        spc_config_from_project(str(tmp_path))
    assert not (tmp_path / "spc" / "config.json").exists()


# ---------------------------------------------------------------------------
# controlplan_recommend_chart — rule table (oracle: test_connector.py) + errors
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("data_type", "n", "kwargs", "expected"),
    [
        ("variable", 1, {}, "I-MR"),
        ("variable", 9, {}, "Xbar-R"),
        ("variable", 10, {}, "Xbar-S"),
        ("variable", 12, {}, "Xbar-S"),  # ceiling boundary — last valid n
        ("attribute", 5, {}, "p"),
        ("attribute", 5, {"defect_based": True, "constant_sample": True}, "c"),
        ("attribute", 5, {"defect_based": True, "constant_sample": False}, "u"),
    ],
)
def test_controlplan_recommend_chart_rule_table(data_type, n, kwargs, expected):
    assert controlplan_recommend_chart(data_type, n, **kwargs) == {
        "recommended_chart": expected
    }


def test_controlplan_recommend_chart_subgroup_size_zero_raises_toolerror():
    with pytest.raises(ToolError, match="subgroup_size"):
        controlplan_recommend_chart("variable", 0)


def test_controlplan_recommend_chart_above_ceiling_raises_toolerror():
    # n=13 is one over the X-bar/S constants ceiling (12) -> structured error.
    with pytest.raises(ToolError, match="exceeds the largest supported"):
        controlplan_recommend_chart("variable", 13)


def test_controlplan_recommend_chart_attribute_has_no_ceiling():
    # Negative guard: the variable-data F-07 ceiling must not leak into attributes.
    assert controlplan_recommend_chart(
        "attribute", 500, defect_based=True, constant_sample=False
    ) == {"recommended_chart": "u"}


# ---------------------------------------------------------------------------
# controlplan_source_index — golden round-trip / empty / invalid / collision
# ---------------------------------------------------------------------------


def test_controlplan_source_index_golden_round_trip():
    model = _relational_model_dict()
    rows = controlplan_build(model)
    index = controlplan_source_index(model)
    # Key set is exactly build's characteristic set (shared traversal).
    assert set(index) == {r["characteristic"] for r in rows}
    # cause_id round-trips to each row's source_cause_id.
    for row in rows:
        assert index[row["characteristic"]]["cause_id"] == row["source_cause_id"]
    # Value shape for the worst-risk cause.
    entry = index["Resin — Uncured"]
    assert set(entry) == {
        "failure_mode_id",
        "cause_id",
        "cause_description",
        "occurrence",
        "component",
    }
    # Worst link (O=8, "Low temperature"), not the O=2 link, per _worst_link.
    assert entry["occurrence"] == 8
    assert entry["cause_description"] == "Low temperature"
    assert entry["component"] == "Resin"


def test_controlplan_source_index_empty_fmea_returns_empty_dict():
    assert controlplan_source_index({"functions": []}) == {}


def test_controlplan_source_index_invalid_model_raises_toolerror():
    with pytest.raises(ToolError, match="validation error"):
        controlplan_source_index({"functions": [{"id": "f1"}]})


def test_controlplan_characteristic_collision_round_trips_through_suffix():
    model = _collision_model_dict()
    rows = controlplan_build(model)
    index = controlplan_source_index(model)
    characteristics = [r["characteristic"] for r in rows]
    # The collision resolves via the ' (F2-M1)' suffix path...
    assert characteristics == [
        "Bracket — Incomplete weld",
        "Bracket — Incomplete weld (F2-M1)",
    ]
    # ...and source_index keys still match build's characteristics one-for-one.
    assert set(index) == set(characteristics)
    for row in rows:
        assert index[row["characteristic"]]["cause_id"] == row["source_cause_id"]


# ===========================================================================
# MSA — Gage R&R (#264). Golden numbers/tolerances are lifted verbatim from
# apps/msa/tests/test_gage_rr_engine.py; the AIAG reference study is loaded
# into list[dict] form (the tool's native input). The MCP layer adds no
# arithmetic, so the boundary must return the already-verified engine numbers.
# ===========================================================================

_AIAG_REFERENCE_STUDY_CSV = (
    Path(__file__).resolve().parents[2] / "msa" / "data" / "aiag_reference_study.csv"
)


def _aiag_study() -> list[dict[str, Any]]:
    """The canonical AIAG 10x3x3 study as the long/tidy list[dict] the tool takes."""
    return load_gage_study_csv(str(_AIAG_REFERENCE_STUDY_CSV)).to_dict("records")


# 6 parts x 3 appraisers x 3 trials with one (P3, appraiser C) cell offset by a
# fixed +1.2 — a genuine part x appraiser interaction. Mirrors
# test_gage_rr_engine.py:1262 (_INDUCED_INTERACTION_DATA), rebuilt as list[dict].
_INDUCED_INTERACTION_STUDY: list[dict[str, Any]] = [
    {
        "part": f"P{p}",
        "appraiser": a,
        "trial": t,
        "measurement": (
            float(p)
            + (0.01 if t == 2 else 0.0)
            + (0.02 if t == 3 else 0.0)
            + (1.2 if (p == 3 and a == "C") else 0.0)
        ),
    }
    for p in range(1, 7)
    for a in ["A", "B", "C"]
    for t in [1, 2, 3]
]

_STUDY_KEYS = ("pev_study", "pav_study", "pgrr_study", "ppv_study")
_TOLERANCE_KEYS = ("pev_tolerance", "pav_tolerance", "pgrr_tolerance", "ppv_tolerance")


# ---------------------------------------------------------------------------
# Golden — Average-and-Range (default method) against the AIAG published form
# ---------------------------------------------------------------------------


def test_msa_gage_rr_average_and_range_golden():
    r = msa_gage_rr(_aiag_study(), tolerance=4.42)
    # Manual-published components (test_gage_rr_engine.py:425-430), rel=1e-2.
    assert r["ev"] == pytest.approx(0.20188, rel=1e-2)
    assert r["av"] == pytest.approx(0.22963, rel=1e-2)
    assert r["grr"] == pytest.approx(0.30576, rel=1e-2)
    assert r["ndc"] == 5
    assert r["verdict"] == "Reject"
    # The tolerance basis reaching the verdict is the #190 fix — pin the number.
    assert r["pgrr_tolerance"] == pytest.approx(41.5067, rel=1e-2)
    # Default method echoed on the payload.
    assert r["method"] == "average_and_range"
    # Study-basis percentages, AIAG Figure III-B 16 (rel=1e-3).
    assert r["pev_study"] == pytest.approx(17.62, rel=1e-3)
    assert r["pav_study"] == pytest.approx(20.04, rel=1e-3)
    assert r["pgrr_study"] == pytest.approx(26.68, rel=1e-3)
    assert r["ppv_study"] == pytest.approx(96.38, rel=1e-3)
    assert r["tv"] == pytest.approx(1.14610, rel=1e-3)
    # Average-and-Range cannot estimate the interaction.
    assert r["interaction"] is None


# ---------------------------------------------------------------------------
# Golden — ANOVA against AIAG Table A 4 / A 5
# ---------------------------------------------------------------------------


def test_msa_gage_rr_anova_golden():
    r = msa_gage_rr(_aiag_study(), method="anova", tolerance=4.42)
    # sigma components, manual 6 dp (test_gage_rr_engine.py:1208-1213), rel=2e-4.
    assert r["ev"] == pytest.approx(0.199933, rel=2e-4)
    assert r["av"] == pytest.approx(0.226838, rel=2e-4)
    assert r["grr"] == pytest.approx(0.302373, rel=2e-4)
    assert r["pv"] == pytest.approx(1.042327, rel=2e-4)
    assert r["tv"] == pytest.approx(1.085, abs=1e-3)
    # F(interaction)=0.434 < F_crit -> pooled to exactly 0, not significant.
    assert r["interaction_f"] == pytest.approx(0.434, abs=1e-3)
    assert r["interaction_significant"] is False
    assert r["interaction"] == 0.0
    assert r["ndc"] == 4
    assert r["method"] == "anova"


# ---------------------------------------------------------------------------
# No-tolerance branch — the four *_tolerance keys null, the four *_study float
# ---------------------------------------------------------------------------


def test_msa_gage_rr_no_tolerance_nulls_tolerance_basis_keeps_study_basis():
    r = msa_gage_rr(_aiag_study())
    for key in _TOLERANCE_KEYS:
        assert r[key] is None, key
    for key in _STUDY_KEYS:
        assert isinstance(r[key], float), key
    # The default method still ran.
    assert r["method"] == "average_and_range"


# ---------------------------------------------------------------------------
# Default method — omitting method= runs Average-and-Range
# ---------------------------------------------------------------------------


def test_msa_gage_rr_default_method_is_average_and_range():
    assert msa_gage_rr(_aiag_study())["method"] == "average_and_range"


# ---------------------------------------------------------------------------
# Interaction-significant branch (ANOVA) — induced interaction diverges upward
# ---------------------------------------------------------------------------


def test_msa_gage_rr_anova_induced_interaction_reports_higher_pgrr():
    anova = msa_gage_rr(_INDUCED_INTERACTION_STUDY, method="anova")
    avg_range = msa_gage_rr(_INDUCED_INTERACTION_STUDY, method="average_and_range")
    # A real interaction — otherwise the divergence claim is vacuous.
    assert anova["interaction_significant"] is True
    assert anova["interaction"] > 0.0
    # ANOVA carries INT^2 into GRR; Average-and-Range cannot see it, so understates.
    assert anova["pgrr_study"] > avg_range["pgrr_study"]
    assert avg_range["interaction"] is None


# ---------------------------------------------------------------------------
# Error paths — every enumerated ValueError branch surfaces as ToolError
# ---------------------------------------------------------------------------

_VALID_ROW = {"part": "P1", "appraiser": "A", "trial": 1, "measurement": 1.0}


def _balanced(parts, appraisers, trials, *, offset=0.0):
    return [
        {
            "part": f"P{p}",
            "appraiser": a,
            "trial": t,
            "measurement": float(p) + (0.02 if a == appraisers[-1] else 0.0) + offset * t,
        }
        for p in range(1, parts + 1)
        for a in appraisers
        for t in range(1, trials + 1)
    ]


def test_msa_gage_rr_empty_study_raises_toolerror():
    with pytest.raises(ToolError, match="at least one measurement"):
        msa_gage_rr([])


def test_msa_gage_rr_missing_required_key_raises_toolerror():
    # changes.md: pd.DataFrame(list_of_dicts) with a missing column surfaces as a
    # ValueError ("Missing required columns"), caught by _call — NOT a raw KeyError.
    # Pinned so a regression to an uncaught KeyError is visible.
    with pytest.raises(ToolError, match="Missing required columns"):
        msa_gage_rr([{"part": "P1"}])


def test_msa_gage_rr_fewer_than_two_parts_raises_toolerror():
    with pytest.raises(ToolError, match="at least 2 parts"):
        msa_gage_rr(_balanced(1, ["A", "B"], 2))


def test_msa_gage_rr_fewer_than_two_appraisers_raises_toolerror():
    with pytest.raises(ToolError, match="at least 2 appraisers"):
        msa_gage_rr(_balanced(2, ["A"], 2))


def test_msa_gage_rr_fewer_than_two_trials_raises_toolerror():
    with pytest.raises(ToolError, match="at least 2 trials"):
        msa_gage_rr(_balanced(2, ["A", "B"], 1))


def test_msa_gage_rr_unbalanced_cells_raise_toolerror():
    study = _balanced(2, ["A", "B"], 2)
    study.append({"part": "P1", "appraiser": "A", "trial": 3, "measurement": 1.5})
    with pytest.raises(ToolError, match="unbalanced"):
        msa_gage_rr(study)


def test_msa_gage_rr_nan_measurement_raises_toolerror():
    study = _balanced(2, ["A", "B"], 2)
    study[0]["measurement"] = float("nan")
    with pytest.raises(ToolError, match="NaN"):
        msa_gage_rr(study)


def test_msa_gage_rr_inf_measurement_raises_toolerror():
    study = _balanced(2, ["A", "B"], 2)
    study[0]["measurement"] = float("inf")
    with pytest.raises(ToolError, match="infinite"):
        msa_gage_rr(study)


def test_msa_gage_rr_nonpositive_tolerance_raises_toolerror():
    with pytest.raises(ToolError, match="positive finite"):
        msa_gage_rr(_balanced(2, ["A", "B"], 2), tolerance=0.0)


def test_msa_gage_rr_unknown_method_raises_toolerror():
    with pytest.raises(ToolError, match="Unknown method: 'bogus'"):
        msa_gage_rr(_balanced(2, ["A", "B"], 2), method="bogus")  # type: ignore[arg-type]
